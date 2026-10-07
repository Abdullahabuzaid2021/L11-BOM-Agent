#!/usr/bin/env python3
"""
Extract BOM data from Excel file and generate JSON for dashboard
"""

import openpyxl
import json
import sys
from pathlib import Path

def extract_bom_data(excel_file):
    """Extract all BOM data from Excel file"""

    print(f"📊 Extracting data from: {excel_file}")

    try:
        wb = openpyxl.load_workbook(excel_file, data_only=True)
        ws = wb['Summary Total']

        # Extract metadata (rows 1-8)
        metadata = {
            'bom_files': [],
            'customers': {},
            'locations': {},
            'focus_ids': {},
            'sfdc_ids': {},
            'target_start_dates': {},
            'target_end_dates': {}
        }

        # Extract components (rows 9+)
        components = []

        # Read BOM file names (row 1, columns G onwards)
        for col in range(7, 25):  # Columns G-X (7-24)
            cell_value = ws.cell(1, col).value
            if cell_value and cell_value != 'Total Quantity':
                bom_file = {
                    'column': col,
                    'name': str(cell_value),
                    'customer': ws.cell(2, col).value or 'Unknown',
                    'location': ws.cell(3, col).value or 'Unknown',
                    'focus_id': ws.cell(5, col).value or 'NOT FOUND',
                    'sfdc_id': ws.cell(6, col).value or 'NOT FOUND',
                    'target_start_date': ws.cell(7, col).value or 'NOT FOUND',
                    'target_end_date': ws.cell(8, col).value or 'NOT FOUND'
                }
                metadata['bom_files'].append(bom_file)

                # Track customers and locations
                if bom_file['customer'] not in metadata['customers']:
                    metadata['customers'][bom_file['customer']] = []
                metadata['customers'][bom_file['customer']].append(col)

                if bom_file['location'] not in metadata['locations']:
                    metadata['locations'][bom_file['location']] = []
                metadata['locations'][bom_file['location']].append(col)

                metadata['focus_ids'][col] = bom_file['focus_id']
                metadata['sfdc_ids'][col] = bom_file['sfdc_id']
                metadata['target_start_dates'][col] = bom_file['target_start_date']
                metadata['target_end_dates'][col] = bom_file['target_end_date']

        # Read component data (rows 9+)
        row = 9
        while True:
            model_pn = ws.cell(row, 1).value
            if not model_pn:
                break

            component = {
                'model_pn': str(model_pn) if model_pn else '',
                'description': ws.cell(row, 2).value or '',
                'nvidia_pn': ws.cell(row, 3).value or '',
                'dell_pn': ws.cell(row, 4).value or '',
                'category': ws.cell(row, 5).value or 'Other',
                'section': ws.cell(row, 6).value or 'Unknown',
                'quantities': {},
                'total_qty': 0
            }

            # Extract quantities for each BOM file
            for col in range(7, 25):
                qty = ws.cell(row, col).value
                if qty and str(qty).isdigit():
                    component['quantities'][col] = int(qty)
                    component['total_qty'] += int(qty)

            if component['quantities']:  # Only add if has quantities
                components.append(component)

            row += 1

        print(f"✅ Extracted {len(metadata['bom_files'])} BOM files")
        print(f"✅ Extracted {len(components)} components")

        return {
            'metadata': metadata,
            'components': components,
            'summary': {
                'total_files': len(metadata['bom_files']),
                'total_components': len(components),
                'customers': list(metadata['customers'].keys()),
                'locations': list(metadata['locations'].keys())
            }
        }

    except Exception as e:
        print(f"❌ Error extracting data: {str(e)}")
        sys.exit(1)

def save_json(data, output_file):
    """Save extracted data as JSON"""

    try:
        with open(output_file, 'w') as f:
            json.dump(data, f, indent=2, default=str)
        print(f"✅ Saved JSON to: {output_file}")
    except Exception as e:
        print(f"❌ Error saving JSON: {str(e)}")
        sys.exit(1)

if __name__ == '__main__':
    excel_file = sys.argv[1] if len(sys.argv) > 1 else 'BOM/BOM_FINAL.xlsx'
    output_file = sys.argv[2] if len(sys.argv) > 2 else 'bom_data.json'

    data = extract_bom_data(excel_file)
    save_json(data, output_file)

    print("\n✅ Data extraction complete!")
    print(f"   Files: {data['summary']['total_files']}")
    print(f"   Components: {data['summary']['total_components']}")
    print(f"   Customers: {', '.join(data['summary']['customers'])}")
