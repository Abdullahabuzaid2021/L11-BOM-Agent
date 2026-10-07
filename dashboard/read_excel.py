import openpyxl

wb = openpyxl.load_workbook('BOM_CONSOLIDATED_FINAL_WITH_SECTIONS.xlsx')
print('Sheets:', wb.sheetnames)
ws = wb.active
print('Active sheet:', ws.title)
print('\nRow 1 values:')
for i, cell in enumerate(ws[1]):
    print(f'Column {i}: {cell.value}')
