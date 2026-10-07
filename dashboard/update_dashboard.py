import openpyxl
from difflib import SequenceMatcher

# IREN FOCUS data from Jira (Focus ID, SFDC ID, Target Start, Target End, attachments)
focus_data = {
    'FOCUS-2586': {
        'sfdc_id': '32103315',
        'target_start': '2026-06-05',
        'target_end': '2026-11-01',
        'attachments': [
            'BOM - NETWORK - 2026-09-11- IREN - 50MW 252 Racks - Sweetwater VR NVL72_SN6600-LD_CORE [Quoted].xlsx',
            'BOM - NETWORK - 2026-09-11- IREN - 50MW 252 Racks - Sweetwater VR NVL72_SN6600-LD_DH [Quoted].xlsx',
        ]
    },
    'FOCUS-2657': {
        'sfdc_id': '31064701',
        'target_start': '2026-01-28',
        'target_end': '2026-11-01',
        'attachments': [
            'BOM - NETWORK - 2026-09-04- IREN -  - VR72 4x Test Racks [Quote].xlsx',
        ]
    },
    'FOCUS-2587': {
        'sfdc_id': '32109457',
        'target_start': '2026-06-05',
        'target_end': '2026-11-01',
        'attachments': [
            'BOM - NETWORK - 2026-09-15- IREN 210 Racks - VR NVL72_SN6600-LD_Core-GOLDEN [Quoted].xlsx',
            'BOM - NETWORK - 2026-09-15- IREN 210 Racks - VR NVL72_SN6600-LD_DH_GOLDEN [Quoted].xlsx',
        ]
    },
}

def similarity(a, b):
    return SequenceMatcher(None, a, b).ratio()

def find_matching_focus(bom_filename, focus_data, threshold=0.7):
    """Find matching Focus ID based on BOM filename similarity"""
    best_match = None
    best_score = 0

    for focus_id, data in focus_data.items():
        for attachment in data['attachments']:
            score = similarity(bom_filename, attachment)
            if score > best_score and score >= threshold:
                best_score = score
                best_match = focus_id

    return best_match, best_score

# Load Excel file
wb = openpyxl.load_workbook('BOM_CONSOLIDATED_FINAL_WITH_SECTIONS.xlsx')
ws = wb.active

# Get BOM filenames from row 1 (columns 6-24)
bom_filenames = []
for col in range(6, 25):
    cell = ws.cell(row=1, column=col)
    if cell.value:
        bom_filenames.append((col, cell.value))

print("Processing BOM filenames...")
for col, filename in bom_filenames:
    print(f"\nColumn {col}: {filename}")
    focus_id, score = find_matching_focus(filename, focus_data)

    if focus_id:
        data = focus_data[focus_id]
        print(f"  Matched with {focus_id} (score: {score:.2f})")

        # Fill in the data
        ws.cell(row=5, column=col, value=focus_id)
        ws.cell(row=6, column=col, value=data['sfdc_id'])
        ws.cell(row=7, column=col, value=data['target_start'])
        ws.cell(row=8, column=col, value=data['target_end'])
    else:
        print(f"  No match found")

# Save the updated file
wb.save('BOM_CONSOLIDATED_FINAL_WITH_SECTIONS_COPY.xlsx')
print("\nFile saved as BOM_CONSOLIDATED_FINAL_WITH_SECTIONS_COPY.xlsx")
