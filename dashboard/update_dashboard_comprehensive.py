import openpyxl
from difflib import SequenceMatcher

# Focus data with NETWORK BOM attachments
# Structure: Focus ID -> {sfdc_id, target_start, target_end, num_bldg, attachments}
focus_data = {
    # IREN
    'FOCUS-2586': {
        'sfdc_id': '32103315',
        'target_start': '2026-06-05',
        'target_end': '2026-11-01',
        'num_bldg': '4 Phases',
        'attachments': [
            'BOM - NETWORK - 2026-09-11- IREN - 50MW 252 Racks - Sweetwater VR NVL72_SN6600-LD_CORE [Quoted].xlsx',
            'BOM - NETWORK - 2026-09-11- IREN - 50MW 252 Racks - Sweetwater VR NVL72_SN6600-LD_DH [Quoted].xlsx',
        ]
    },
    'FOCUS-2657': {
        'sfdc_id': '31064701',
        'target_start': '2026-01-28',
        'target_end': '2026-11-01',
        'num_bldg': '1',
        'attachments': [
            'BOM - NETWORK - 2026-09-04- IREN -  - VR72 4x Test Racks [Quote].xlsx',
        ]
    },
    'FOCUS-2587': {
        'sfdc_id': '32109457',
        'target_start': '2026-06-05',
        'target_end': '2026-11-01',
        'num_bldg': '2 Phases',
        'attachments': [
            'BOM - NETWORK - 2026-09-15- IREN 210 Racks - VR NVL72_SN6600-LD_Core-GOLDEN [Quoted].xlsx',
            'BOM - NETWORK - 2026-09-15- IREN 210 Racks - VR NVL72_SN6600-LD_DH_GOLDEN [Quoted].xlsx',
        ]
    },
    # NScale
    'FOCUS-704': {
        'sfdc_id': None,
        'target_start': '2026-02-02',
        'target_end': '2026-11-01',
        'num_bldg': None,
        'attachments': [
            'CTO BOM Nscale VR Texas 16k 2026.6.30.xlsx',
        ]
    },
    'FOCUS-694': {
        'sfdc_id': None,
        'target_start': '2026-02-02',
        'target_end': '2026-11-01',
        'num_bldg': None,
        'attachments': [
            'CTO BOM Nscale VR London 14k 2026.6.30.xlsx',
        ]
    },
    'FOCUS-1458': {
        'sfdc_id': None,
        'target_start': None,
        'target_end': None,
        'num_bldg': None,
        'attachments': [
            'CTO BOM NScale VR Portugal SIN02 5k 6k 8k 12k 2026.9.15.xlsx',
        ]
    },
    'FOCUS-684': {
        'sfdc_id': None,
        'target_start': None,
        'target_end': None,
        'num_bldg': None,
        'attachments': [
            'CTO BOM Nscale VR Norway Kvanndal South 15k 16k 2026.5.8.xlsx',
        ]
    },
    'FOCUS-1468': {
        'sfdc_id': None,
        'target_start': None,
        'target_end': None,
        'num_bldg': None,
        'attachments': [
            'CTO BOM Nscale VR PS 1k 2026.5.8.xlsx',
            'CTO BOM Nscale VR PS 1k 2026.9.15 (Semi LLD).xlsx',
        ]
    },
    'FOCUS-1364': {
        'sfdc_id': None,
        'target_start': '2027-03-31',
        'target_end': '2027-03-31',
        'num_bldg': '3 Clusters',
        'attachments': [
            'CTO BOM NScale VR Norway Narvik North 8k 17k 2026.9.15.xlsx',
            'NScale - NETWORK BOM - Narvik-8k.xlsx',
            'NScale - NETWORK BOM - Narvik-17k.xlsx',
            'NScale - NETWORK BOM - Narvik-8k-6810.xlsx',
        ]
    },
    'FOCUS-2126': {
        'sfdc_id': None,
        'target_start': None,
        'target_end': None,
        'attachments': [
            'CTO BOM NScale VR Figure RoCE-Shfl-8P 5k+5k 2026.8.30.xlsx',
        ]
    },
    'FOCUS-2118': {
        'sfdc_id': None,
        'target_start': None,
        'target_end': None,
        'attachments': [
            'CTO BOM NScale WVA VR RoCE-Shfl-8P 144k 4x36k 2026.9.21.xlsx',
        ]
    },
    'FOCUS-1448': {
        'sfdc_id': None,
        'target_start': None,
        'target_end': None,
        'attachments': [
            'CTO BOM Nscale GB300 London Glass 4.6k 2026.5.11.xlsx',
        ]
    },
    'FOCUS-3021': {
        'sfdc_id': None,
        'target_start': None,
        'target_end': None,
        'attachments': [
            'CTO BOM FINLAND EuroHPC VR RoCE-CPO-8P 7SU 2026.9.21.xlsx',
        ]
    },
    # Anthropic
    'FOCUS-2116': {
        'sfdc_id': None,
        'target_start': '2027-02-07',
        'target_end': '2027-07-31',
        'num_bldg': '3 Locations',
        'attachments': [
            'PNL-NETWORK BOM-2026-09-23-Anthropic-Australia_Q1_Q2-21k GPUs VR NVL72.xlsx',
            'PNL-NETWORK BOM-2026-09-23-Anthropic-Canada_Q2-32k GPUs VR NVL72.xlsx',
            'PNL-NETWORK BOM-2026-09-24-Anthropic-US-Q1-Q2-Cluster1-129k GPUs VR NVL72.xlsx',
            'PNL-NETWORK BOM-2026-09-24-Anthropic-US-Q1-Q2-Cluster2-129k GPUs VR NVL72.xlsx',
            'PNL-NETWORK BOM-2026-09-25-Anthropic-US_Q1_Q2-Cluster3-9k GPUs VR NVL72.xlsx',
        ]
    },
}

def similarity(a, b):
    return SequenceMatcher(None, a, b).ratio()

def find_matching_focus(bom_filename, focus_data, threshold=0.7):
    """Find matching Focus ID based on BOM filename similarity"""
    best_match = None
    best_score = 0

    for focus_id, data in focus_data.items():
        for attachment in data['attachments']:
            score = similarity(bom_filename, attachment)
            if score > best_score and score >= threshold:
                best_score = score
                best_match = focus_id

    return best_match, best_score

# Load Excel file
wb = openpyxl.load_workbook('BOM_CONSOLIDATED_FINAL_WITH_SECTIONS.xlsx')
ws = wb.active

# Get BOM filenames from row 1 (columns 6-24)
bom_filenames = []
for col in range(6, 25):
    cell = ws.cell(row=1, column=col)
    if cell.value and 'NETWORK' in str(cell.value).upper():
        bom_filenames.append((col, cell.value))

print("Processing NETWORK BOM filenames...")
matches_found = 0
for col, filename in bom_filenames:
    print(f"\nColumn {col}: {filename}")
    focus_id, score = find_matching_focus(filename, focus_data)

    if focus_id:
        data = focus_data[focus_id]
        print(f"  Matched with {focus_id} (score: {score:.2f})")
        matches_found += 1

        # Fill in the data
        ws.cell(row=4, column=col, value=data.get('num_bldg'))
        ws.cell(row=5, column=col, value=focus_id)
        ws.cell(row=6, column=col, value=data['sfdc_id'])
        ws.cell(row=7, column=col, value=data['target_start'])
        ws.cell(row=8, column=col, value=data['target_end'])
    else:
        print(f"  No match found")

# Save the updated file
wb.save('BOM_CONSOLIDATED_FINAL_WITH_SECTIONS_COPY.xlsx')
print(f"\nFile saved as BOM_CONSOLIDATED_FINAL_WITH_SECTIONS_COPY.xlsx")
print(f"Total matches found: {matches_found}")
