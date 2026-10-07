import openpyxl

wb = openpyxl.load_workbook('BOM_CONSOLIDATED_FINAL_WITH_SECTIONS_COPY.xlsx')
ws = wb.active

print('Verifying updated data in rows 4-8:')
print('\nRow 4 (# of Bldg):')
for col in range(6, 25):
    cell = ws.cell(row=4, column=col)
    print(f'  Column {col}: {cell.value}')

print('\nRow 5 (Focus ID):')
for col in range(6, 25):
    cell = ws.cell(row=5, column=col)
    print(f'  Column {col}: {cell.value}')

print('\nRow 6 (SFDC ID):')
for col in range(6, 25):
    cell = ws.cell(row=6, column=col)
    print(f'  Column {col}: {cell.value}')

print('\nRow 7 (Target Start Date):')
for col in range(6, 25):
    cell = ws.cell(row=7, column=col)
    print(f'  Column {col}: {cell.value}')

print('\nRow 8 (Target End Date):')
for col in range(6, 25):
    cell = ws.cell(row=8, column=col)
    print(f'  Column {col}: {cell.value}')
